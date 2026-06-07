from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = Path(__file__).resolve().parent.parent / "docs" / "midterm" / "MIDTERM_QA.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Malgun Gothic"
style.font.size = Pt(11)


def add_qa(num, question, answer):
    p = doc.add_paragraph()
    q = p.add_run(f"Q{num}. {question}")
    q.bold = True
    q.font.size = Pt(11)
    doc.add_paragraph(f"A: {answer}")


def add_section(title, level=1):
    doc.add_heading(title, level=level)


doc.add_heading("IoT Midterm Q&A", level=0)
sub = doc.add_paragraph(
    "Motion-Adaptive AI Companion Robot\n"
    "LLM-based Emotional Interactive AIoT System"
)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Team: Park Hyodam (202304709, Captain) · Jo Hyeonjin (202304438)")
doc.add_paragraph("")

# --- Korean ---
add_section("Part 1 — 한국어 Q&A", 1)
add_section("교수님 예상 질문", 2)

ko_prof = [
    ("카메라와 초음파 센서를 왜 둘 다 쓰나요? 하나만으로는 안 되나요?",
     "카메라는 손을 흔드는 것 같은 motion을 감지하고, 초음파는 20cm 이하로 가까이 오는 proximity를 감지합니다. 둘을 함께 쓰면 사용자 상황을 더 잘 포착할 수 있어서 motion-adaptive 목표에 더 맞습니다."),
    ("idle, approach, stay, alert 상태의 차이는 무엇인가요?",
     "Idle은 대기, Approach는 20cm 이하 접근, Stay는 근처에 머무르며 대화 유도, Alert는 갑작스러운 큰 움직임입니다. 현재는 idle/motion/reply까지 구현했고 최종에 세분화할 예정입니다."),
    ("키워드만 쓰는데 LLM 기반 프로젝트가 맞나요?",
     "중간 발표 시점에서는 keyword matching이 기본입니다. 네트워크가 되면 Gemini API로 보강하고, Pi 로컬 LLM은 단계적으로 적용합니다."),
    ("버저는 왜 GPIO 27번인가요?",
     "초음파 ECHO 핀이 GPIO 18을 사용하기 때문에 버저를 27번으로 분리했습니다."),
    ("스피커가 없는데 사용자는 어떻게 로봇의 답을 확인하나요?",
     "답변은 LCD 텍스트로 표시합니다. 최종에는 voice-assistant 참고자료로 음성 확장이 가능합니다."),
    ("detect_motion()은 어떻게 동작하나요?",
     "2초마다 rpicam-still로 촬영 후 중앙 ROI grayscale 차이 평균이 threshold 25를 넘으면 motion으로 판단합니다."),
    ("show_motion()에서는 무엇이 일어나나요?",
     "LED ON, LCD 인사 문구, 버저 0.4초입니다. led_lcd.txt 초안에서 버저까지 통합했습니다."),
    ("지금까지 실제로 무엇을 완료했나요?",
     "LCD·LED·버저, 초음파, OpenCV motion, idle/motion/reply 상태 완료. companion.py 통합 테스트 진행 중. Lab 9도 완료."),
    ("최종 제출 전에 무엇이 남았나요?",
     "통합 안정화, approach/stay/alert 정교화, 감정형 응답, 오류 처리, 데모 영상·보고서."),
    ("팀원 역할은 어떻게 나눴나요?",
     "조현진(A): motion, reply, companion. 박효담(B): lcd_actuator, 하드웨어, 보고서 취합."),
]
for i, (q, a) in enumerate(ko_prof, 1):
    add_qa(i, q, a)

add_section("학생 예상 질문", 2)
ko_stu = [
    ("인터넷 없이도 동작하나요?", "motion, LCD, LED, 버저는 오프라인 동작합니다."),
    ("얼굴이나 감정 인식도 하나요?", "아직은 아닙니다. motion, distance, keyword 기반입니다."),
    ("초음파 센서가 -1을 반환하면?", "센서 오류로 보고 camera motion만 사용하거나 idle 메시지를 표시합니다."),
    ("Lab 9와 무엇이 다른가요?", "Lab 9는 터치→LED. 기말은 카메라+초음파+LCD+대화 AIoT 시스템입니다."),
]
for i, (q, a) in enumerate(ko_stu, 11):
    add_qa(i, q, a)

doc.add_paragraph("")
p = doc.add_paragraph()
r = p.add_run("모르는 질문: ")
r.bold = True
doc.add_paragraph(
    "좋은 질문입니다. 아직 완전히 구현하지는 못했지만, 최종 버전 계획에 포함되어 있습니다."
)

doc.add_page_break()

# --- English ---
add_section("Part 2 — English Q&A", 1)
add_section("Professor — Likely Questions", 2)

en_prof = [
    ("Why do you use both a camera and an ultrasonic sensor?",
     "The camera detects motion; the ultrasonic sensor detects proximity within 20 cm. Together they cover more user situations."),
    ("What is the difference between idle, approach, stay, and alert?",
     "Idle: no user. Approach: within 20 cm. Stay: user remains nearby. Alert: sudden large motion. We plan to refine this flow."),
    ("Is this really LLM-based if you only use keyword matching?",
     "Keyword matching is our main method now; we optionally use Gemini API when online."),
    ("Why is the buzzer on GPIO 27?",
     "GPIO 18 is used by ultrasonic ECHO, so we moved the buzzer to GPIO 27."),
    ("Why is there no speaker?",
     "Replies are shown as text on the LCD. Voice may be added later."),
    ("How does detect_motion() work?",
     "Every 2 seconds we capture a frame and compare center ROI grayscale diff; above threshold 25 means motion."),
    ("What happens in show_motion()?",
     "LED on, greeting on LCD, buzzer beeps 0.4 seconds."),
    ("What have you completed so far?",
     "LCD/LED/buzzer, ultrasonic, OpenCV motion, basic states. Integration testing in progress."),
    ("What remains before final submission?",
     "Integration, state refinement, emotional replies, error handling, demo video and report."),
    ("How are tasks divided?",
     "Jo Hyeonjin (A): motion, reply, companion. Park Hyodam (B): lcd_actuator, hardware, report."),
]
for i, (q, a) in enumerate(en_prof, 1):
    add_qa(i, q, a)

add_section("Student — Likely Questions", 2)
en_stu = [
    ("Does it work without internet?", "Motion, LCD, LED, and buzzer work offline."),
    ("Can it recognize faces or emotions?", "Not yet. Motion, distance, and keywords for now."),
    ("What if ultrasonic returns -1?", "Treat as sensor error; fall back to camera motion."),
    ("How is this different from Lab 9?", "Lab 9 was touch to LED; ours is a full AIoT system."),
]
for i, (q, a) in enumerate(en_stu, 11):
    add_qa(i, q, a)

doc.add_paragraph("")
p = doc.add_paragraph()
r = p.add_run("If you don't know the answer: ")
r.bold = True
doc.add_paragraph(
    "That is a good question. We have considered it, but we have not fully implemented it yet. "
    "It is part of our plan for the final version."
)

add_section("Top 3 to Memorize", 2)
for line in [
    "1. Why camera + ultrasonic? — Motion and proximity detect different situations.",
    "2. Why buzzer on GPIO 27? — GPIO 18 is used by ultrasonic ECHO.",
    "3. Done vs. planned? — Sensors/actuators work; refining states and replies.",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.save(OUT)
print("Saved:", OUT)
