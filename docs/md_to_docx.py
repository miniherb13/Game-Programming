#!/usr/bin/env python3
"""Convert project markdown report to .docx (python-docx)."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_east_asia_font(run, name: str = "Malgun Gothic") -> None:
    run.font.name = name
    r = run._element.rPr
    if r is not None:
        r.rFonts.set(qn("w:eastAsia"), name)


def add_formatted_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
            set_east_asia_font(run)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        else:
            run = p.add_run(part)
            set_east_asia_font(run)


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def is_table_separator(line: str) -> bool:
    cells = parse_table_row(line)
    return all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells if c)


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15

    i = 0
    in_code = False
    code_lines: list[str] = []
    in_blockquote = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                pf = p.paragraph_format
                pf.left_indent = Cm(0.5)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("─" * 40)
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_heading(stripped[2:].strip(), level=0)
            for run in p.runs:
                set_east_asia_font(run)
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_heading(stripped[3:].strip(), level=1)
            for run in p.runs:
                set_east_asia_font(run)
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_heading(stripped[4:].strip(), level=2)
            for run in p.runs:
                set_east_asia_font(run)
            i += 1
            continue

        if stripped.startswith("#### "):
            p = doc.add_heading(stripped[5:].strip(), level=3)
            for run in p.runs:
                set_east_asia_font(run)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            headers = parse_table_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            col_count = len(headers)
            table = doc.add_table(rows=1 + len(rows), cols=col_count)
            table.style = "Table Grid"
            for c, h in enumerate(headers):
                cell = table.rows[0].cells[c]
                cell.text = h
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    set_east_asia_font(run)
            for r, row in enumerate(rows):
                for c in range(col_count):
                    text = row[c] if c < len(row) else ""
                    cell = table.rows[r + 1].cells[c]
                    cell.text = text
                    for run in cell.paragraphs[0].runs:
                        set_east_asia_font(run)
            doc.add_paragraph()
            continue

        if stripped.startswith("> "):
            add_formatted_paragraph(doc, stripped[2:], style=None)
            i += 1
            continue

        if stripped.startswith("- [ ]") or stripped.startswith("- [x]"):
            mark = "☑" if stripped.startswith("- [x]") else "☐"
            add_formatted_paragraph(doc, f"{mark} {stripped[5:].strip()}", style="List Bullet")
            i += 1
            continue

        if stripped.startswith("- "):
            add_formatted_paragraph(doc, stripped[2:], style="List Bullet")
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            text = re.sub(r"^\d+\.\s", "", stripped)
            add_formatted_paragraph(doc, text, style="List Number")
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        add_formatted_paragraph(doc, stripped)
        i += 1

    doc.save(str(docx_path))
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    root = Path(__file__).resolve().parent
    md = root / "final-report-draft.md"
    out = root / "final-report-draft.docx"
    if len(sys.argv) >= 2:
        md = Path(sys.argv[1])
    if len(sys.argv) >= 3:
        out = Path(sys.argv[2])
    convert(md, out)
